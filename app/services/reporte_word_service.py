import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.config import Config


# =========================================================
# PALETA VISUAL
# =========================================================

AZUL_NOCHE = "102A43"
AZUL_PRINCIPAL = "1F4E78"
AZUL_MEDIO = "3E7CB1"
AZUL_CLARO = "EAF2F8"

GRIS_FONDO = "F7F9FC"
GRIS_BORDE = "D5DEE8"
GRIS_TEXTO = RGBColor(82, 96, 109)

VERDE = "2EAD67"
NARANJA = "F59E42"
ROJO = "D9534F"


# =========================================================
# FUNCIONES AUXILIARES
# =========================================================

def _texto(valor, vacio="Sin información"):
    if valor is None or str(valor).strip() == "":
        return vacio

    return str(valor)


def _fecha_legible(fecha):
    if not fecha:
        return "Sin fecha"

    try:
        return fecha.strftime("%d/%m/%Y")
    except AttributeError:
        return str(fecha)


def _estado_legible(estado):
    estados = {
        "CUMPLE": "Cumple",
        "PARCIAL": "Cumplimiento parcial",
        "NO_CUMPLE": "No cumple",
    }

    return estados.get(estado, estado)


def _color_estado(estado):
    if estado == "CUMPLE":
        return VERDE

    if estado == "PARCIAL":
        return NARANJA

    if estado == "NO_CUMPLE":
        return ROJO

    return AZUL_PRINCIPAL


def _colorear_celda(celda, color):
    """
    Aplica color de fondo a una celda de Word.
    """

    tc_pr = celda._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), color)


def _margen_celda(celda, top=100, start=100, bottom=100, end=100):
    """
    Agrega espacio interno a las celdas.
    """

    tc = celda._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margen, valor in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():

        nodo = tc_mar.find(qn(f"w:{margen}"))

        if nodo is None:
            nodo = OxmlElement(f"w:{margen}")
            tc_mar.append(nodo)

        nodo.set(qn("w:w"), str(valor))
        nodo.set(qn("w:type"), "dxa")


def _formatear_encabezado_tabla(fila):
    """
    Aplica el estilo azul a la primera fila de una tabla.
    """

    for celda in fila.cells:
        _colorear_celda(celda, AZUL_PRINCIPAL)
        _margen_celda(celda)

        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for parrafo in celda.paragraphs:
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in parrafo.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)


def _titulo_seccion(doc, texto):
    """
    Crea títulos uniformes para las secciones.
    """

    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(7)

    run = p.add_run(texto.upper())

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(31, 78, 120)

    return p


def _generar_conclusiones(consolidado):
    t = consolidado["totales"]

    conclusiones = []

    total = t["total_registros"]

    if total == 0:
        return [
            "No se registraron actividades SG-SST durante el periodo seleccionado."
        ]

    conclusiones.append(
        f"Durante el periodo se evaluaron {total} actividades SG-SST. "
        f"El cumplimiento general alcanzado fue del "
        f"{t['porcentaje_avance']}%."
    )

    conclusiones.append(
        f"De las {total} actividades evaluadas, "
        f"{t['cumplidos']} alcanzaron cumplimiento total, "
        f"{t['parciales']} presentaron cumplimiento parcial y "
        f"{t['no_cumplidos']} fueron clasificadas como no cumplidas."
    )

    if t["parciales"] > 0:
        conclusiones.append(
            f"Se identificaron {t['parciales']} actividades con "
            f"cumplimiento parcial que requieren seguimiento para "
            f"verificar su cierre."
        )

    if t["no_cumplidos"] > 0:
        conclusiones.append(
            f"Se identificaron {t['no_cumplidos']} actividades en estado "
            f"No Cumple que requieren atención y seguimiento."
        )

    fases = consolidado.get("por_fase", [])

    if fases:
        fase_menor = min(
            fases,
            key=lambda x: x["porcentaje_avance"]
        )

        conclusiones.append(
            f"La fase PHVA con menor nivel de cumplimiento fue "
            f"{fase_menor['nombre_fase']}, con un avance del "
            f"{fase_menor['porcentaje_avance']}%."
        )

    sedes = consolidado.get("por_sede", [])

    if len(sedes) > 1:
        mejor = max(
            sedes,
            key=lambda x: x["porcentaje_avance"]
        )

        menor = min(
            sedes,
            key=lambda x: x["porcentaje_avance"]
        )

        conclusiones.append(
            f"La sede con mayor cumplimiento fue "
            f"{mejor['nombre_sede']}, con "
            f"{mejor['porcentaje_avance']}%."
        )

        conclusiones.append(
            f"La sede con menor cumplimiento fue "
            f"{menor['nombre_sede']}, con "
            f"{menor['porcentaje_avance']}%, por lo que representa "
            f"un punto prioritario de seguimiento."
        )

    return conclusiones


# =========================================================
# GENERACIÓN DEL WORD
# =========================================================

def generar_reporte_word(
    consolidado: dict,
    generado_por: str,
    nombre_archivo: str
) -> str:

    doc = Document()

    # =====================================================
    # CONFIGURACIÓN DE PÁGINA
    # =====================================================

    seccion = doc.sections[0]

    # Horizontal
    seccion.orientation = WD_ORIENT.LANDSCAPE

    seccion.page_width = Cm(27.94)
    seccion.page_height = Cm(21.59)

    seccion.top_margin = Cm(1.5)
    seccion.bottom_margin = Cm(1.5)
    seccion.left_margin = Cm(1.5)
    seccion.right_margin = Cm(1.5)

    # Fuente general
    estilo_normal = doc.styles["Normal"]

    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(9)
    estilo_normal.font.color.rgb = RGBColor(36, 59, 83)

    # =====================================================
    # ENCABEZADO SUPERIOR
    # =====================================================

    encabezado = seccion.header

    tabla_header = encabezado.add_table(
        rows=1,
        cols=1,
        width=Cm(24.5)
    )

    tabla_header.alignment = WD_TABLE_ALIGNMENT.CENTER

    celda_header = tabla_header.cell(0, 0)

    _colorear_celda(
        celda_header,
        AZUL_NOCHE
    )

    p_header = celda_header.paragraphs[0]

    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_header = p_header.add_run(
        "SISTEMA DE GESTIÓN DE SEGURIDAD Y SALUD EN EL TRABAJO"
    )

    run_header.bold = True
    run_header.font.color.rgb = RGBColor(255, 255, 255)
    run_header.font.size = Pt(8)

    # =====================================================
    # TÍTULO
    # =====================================================

    titulo = doc.add_paragraph()

    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(3)

    run = titulo.add_run(
        "INFORME DE GESTIÓN SG-SST"
    )

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(16, 42, 67)

    subtitulo = doc.add_paragraph()

    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo.paragraph_format.space_after = Pt(10)

    run = subtitulo.add_run(
        "Sistema de Gestión de Seguridad y Salud en el Trabajo"
    )

    run.font.size = Pt(9)
    run.font.color.rgb = GRIS_TEXTO

    # =====================================================
    # INFORMACIÓN GENERAL
    # =====================================================

    tabla_info = doc.add_table(
        rows=1,
        cols=3
    )

    tabla_info.alignment = WD_TABLE_ALIGNMENT.CENTER

    datos_info = [
        (
            "TIPO DE CORTE",
            consolidado["tipo_corte"]
        ),
        (
            "PERIODO EVALUADO",
            consolidado["periodo"]
        ),
        (
            "GENERADO POR",
            generado_por
        ),
    ]

    for i, (titulo_info, valor) in enumerate(datos_info):

        celda = tabla_info.cell(0, i)

        _colorear_celda(
            celda,
            AZUL_CLARO
        )

        _margen_celda(celda)

        p = celda.paragraphs[0]

        run = p.add_run(
            f"{titulo_info}\n"
        )

        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(31, 78, 120)

        run2 = p.add_run(
            str(valor)
        )

        run2.font.size = Pt(9)

    # =====================================================
    # 1. RESUMEN EJECUTIVO
    # =====================================================

    _titulo_seccion(
        doc,
        "1. Resumen ejecutivo"
    )

    t = consolidado["totales"]

    p = doc.add_paragraph()

    p.paragraph_format.space_after = Pt(8)

    p.add_run(
        f"Durante el periodo {consolidado['periodo']} se evaluaron "
    )

    run = p.add_run(
        f"{t['total_registros']} actividades"
    )

    run.bold = True

    p.add_run(
        " del Sistema de Gestión de Seguridad y Salud en el Trabajo. "
        f"De estas, {t['cumplidos']} alcanzaron cumplimiento total, "
        f"{t['parciales']} presentaron cumplimiento parcial y "
        f"{t['no_cumplidos']} fueron clasificadas como no cumplidas. "
        "El nivel general de cumplimiento del periodo fue del "
    )

    run = p.add_run(
        f"{t['porcentaje_avance']}%."
    )

    run.bold = True

    # =====================================================
    # INDICADORES
    # =====================================================

    indicadores = doc.add_table(
        rows=2,
        cols=5
    )

    indicadores.alignment = WD_TABLE_ALIGNMENT.CENTER

    encabezados = [
        "ACTIVIDADES EVALUADAS",
        "CUMPLIDAS",
        "PARCIALES",
        "NO CUMPLIDAS",
        "CUMPLIMIENTO GENERAL",
    ]

    valores = [
        t["total_registros"],
        t["cumplidos"],
        t["parciales"],
        t["no_cumplidos"],
        f"{t['porcentaje_avance']}%",
    ]

    for i, texto in enumerate(encabezados):

        celda = indicadores.cell(0, i)

        celda.text = texto

    _formatear_encabezado_tabla(
        indicadores.rows[0]
    )

    for i, valor in enumerate(valores):

        celda = indicadores.cell(1, i)

        _margen_celda(celda)

        p = celda.paragraphs[0]

        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run(
            str(valor)
        )

        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(16, 42, 67)

    # =====================================================
    # 2. DESEMPEÑO POR SEDE
    # =====================================================

    _titulo_seccion(
        doc,
        "2. Desempeño por sede"
    )

    sedes = consolidado.get(
        "por_sede",
        []
    )

    if sedes:

        tabla_sedes = doc.add_table(
            rows=1,
            cols=6
        )

        tabla_sedes.alignment = WD_TABLE_ALIGNMENT.CENTER

        encabezados = [
            "Sede",
            "Evaluadas",
            "Cumple",
            "Parcial",
            "No cumple",
            "% Cumplimiento",
        ]

        for i, texto in enumerate(encabezados):
            tabla_sedes.cell(0, i).text = texto

        _formatear_encabezado_tabla(
            tabla_sedes.rows[0]
        )

        for sede in sedes:

            fila = tabla_sedes.add_row().cells

            valores_sede = [
                sede["nombre_sede"],
                sede["total"],
                sede["cumplidos"],
                sede["parciales"],
                sede["no_cumplidos"],
                f"{sede['porcentaje_avance']}%",
            ]

            for i, valor in enumerate(valores_sede):

                fila[i].text = str(valor)

                _margen_celda(
                    fila[i]
                )

                fila[i].vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

                if i > 0:
                    fila[i].paragraphs[0].alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                    )

        if len(sedes) > 1:

            mejor = max(
                sedes,
                key=lambda x: x["porcentaje_avance"]
            )

            menor = min(
                sedes,
                key=lambda x: x["porcentaje_avance"]
            )

            p = doc.add_paragraph()

            p.paragraph_format.space_before = Pt(7)

            run = p.add_run(
                "Lectura del resultado: "
            )

            run.bold = True

            p.add_run(
                f"{mejor['nombre_sede']} presentó el mayor nivel "
                f"de cumplimiento ({mejor['porcentaje_avance']}%). "
                f"{menor['nombre_sede']} registró el menor resultado "
                f"del periodo ({menor['porcentaje_avance']}%)."
            )

    else:
        doc.add_paragraph(
            "No existen registros por sede para el periodo seleccionado."
        )

    # =====================================================
    # 3. ANÁLISIS POR FASE PHVA
    # =====================================================

    _titulo_seccion(
        doc,
        "3. Análisis por fase PHVA"
    )

    tabla_fase = doc.add_table(
        rows=1,
        cols=6
    )

    tabla_fase.alignment = WD_TABLE_ALIGNMENT.CENTER

    encabezados = [
        "Fase",
        "Actividades",
        "Cumple",
        "Parcial",
        "No cumple",
        "% Avance",
    ]

    for i, texto in enumerate(encabezados):
        tabla_fase.cell(0, i).text = texto

    _formatear_encabezado_tabla(
        tabla_fase.rows[0]
    )

    for fase in consolidado["por_fase"]:

        fila = tabla_fase.add_row().cells

        valores_fase = [
            fase["nombre_fase"],
            fase["total"],
            fase["cumplidos"],
            fase["parciales"],
            fase["no_cumplidos"],
            f"{fase['porcentaje_avance']}%",
        ]

        for i, valor in enumerate(valores_fase):

            fila[i].text = str(valor)

            _margen_celda(fila[i])

            if i > 0:
                fila[i].paragraphs[0].alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

    for fase in consolidado["por_fase"]:

        p = doc.add_paragraph()

        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)

        run = p.add_run(
            f"{fase['nombre_fase']}: "
        )

        run.bold = True

        p.add_run(
            f"se evaluaron {fase['total']} actividades; "
            f"{fase['cumplidos']} cumplieron completamente, "
            f"{fase['parciales']} presentaron cumplimiento parcial "
            f"y {fase['no_cumplidos']} no cumplieron. "
            f"El avance de la fase fue del "
            f"{fase['porcentaje_avance']}%."
        )

    # =====================================================
    # SALTO DE PÁGINA
    # =====================================================

    doc.add_page_break()

    # =====================================================
    # 4. ACTIVIDADES REALIZADAS
    # =====================================================

    _titulo_seccion(
        doc,
        "4. Actividades realizadas durante el periodo"
    )

    doc.add_paragraph(
        "A continuación se presentan las actividades registradas, "
        "indicando la sede donde fueron realizadas, su responsable, "
        "fecha, resultado, evidencia requerida y observaciones."
    )

    for item in consolidado["detalle"]:

        tabla = doc.add_table(
            rows=8,
            cols=2
        )

        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

        datos = [
            (
                f"ACTIVIDAD {item['codigo_item']}",
                _texto(
                    item["descripcion_item"]
                )
            ),
            (
                "Fase PHVA",
                _texto(
                    item["nombre_fase"]
                )
            ),
            (
                "Sede",
                _texto(
                    item["nombre_sede"]
                )
            ),
            (
                "Responsable",
                _texto(
                    item["responsable"]
                )
            ),
            (
                "Fecha",
                _fecha_legible(
                    item["fecha_registro"]
                )
            ),
            (
                "Resultado",
                f"{_estado_legible(item['estado_cumplimiento'])} "
                f"— {item['porcentaje_item']}%"
            ),
            (
                "Evidencia requerida",
                _texto(
                    item["evidencia_requerida"]
                )
            ),
            (
                "Observación",
                _texto(
                    item["observacion"],
                    "Sin observación"
                )
            ),
        ]

        for indice, (etiqueta, valor) in enumerate(datos):

            celda_etiqueta = tabla.cell(
                indice,
                0
            )

            celda_valor = tabla.cell(
                indice,
                1
            )

            celda_etiqueta.text = etiqueta
            celda_valor.text = str(valor)

            _margen_celda(celda_etiqueta)
            _margen_celda(celda_valor)

            if indice == 0:

                _colorear_celda(
                    celda_etiqueta,
                    AZUL_CLARO
                )

                _colorear_celda(
                    celda_valor,
                    AZUL_CLARO
                )

                for celda in (
                    celda_etiqueta,
                    celda_valor
                ):
                    for run in celda.paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(
                            16,
                            42,
                            67
                        )

            else:
                _colorear_celda(
                    celda_etiqueta,
                    GRIS_FONDO
                )

                for run in celda_etiqueta.paragraphs[0].runs:
                    run.bold = True

            # Resultado con color según estado
            if indice == 5:

                color = _color_estado(
                    item["estado_cumplimiento"]
                )

                _colorear_celda(
                    celda_valor,
                    color
                )

                for run in celda_valor.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(
                        255,
                        255,
                        255
                    )

        doc.add_paragraph()

    # =====================================================
    # 5. ACTIVIDADES QUE REQUIEREN SEGUIMIENTO
    # =====================================================

    doc.add_page_break()

    _titulo_seccion(
        doc,
        "5. Actividades que requieren seguimiento"
    )

    pendientes = [
        item
        for item in consolidado["detalle"]
        if item["estado_cumplimiento"] in (
            "PARCIAL",
            "NO_CUMPLE"
        )
    ]

    if pendientes:

        tabla_pendientes = doc.add_table(
            rows=1,
            cols=6
        )

        tabla_pendientes.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        encabezados = [
            "Actividad",
            "Sede",
            "Responsable",
            "Estado",
            "Fecha",
            "Observación",
        ]

        for i, texto in enumerate(encabezados):
            tabla_pendientes.cell(
                0,
                i
            ).text = texto

        _formatear_encabezado_tabla(
            tabla_pendientes.rows[0]
        )

        for item in pendientes:

            fila = tabla_pendientes.add_row().cells

            valores = [
                f"{item['codigo_item']} - "
                f"{_texto(item['descripcion_item'])}",
                _texto(item["nombre_sede"]),
                _texto(item["responsable"]),
                _estado_legible(
                    item["estado_cumplimiento"]
                ),
                _fecha_legible(
                    item["fecha_registro"]
                ),
                _texto(
                    item["observacion"],
                    "Sin observación"
                ),
            ]

            for i, valor in enumerate(valores):

                fila[i].text = str(valor)

                _margen_celda(fila[i])

            # Colorear solamente el estado
            color = _color_estado(
                item["estado_cumplimiento"]
            )

            _colorear_celda(
                fila[3],
                color
            )

            for run in fila[3].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(
                    255,
                    255,
                    255
                )

    else:

        doc.add_paragraph(
            "Durante el periodo no se registraron actividades "
            "con cumplimiento parcial o estado No Cumple."
        )

    # =====================================================
    # 6. CONCLUSIONES
    # =====================================================

    _titulo_seccion(
        doc,
        "6. Conclusiones del periodo"
    )

    conclusiones = _generar_conclusiones(
        consolidado
    )

    for numero, conclusion in enumerate(
        conclusiones,
        start=1
    ):

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(4)

        run = p.add_run(
            f"{numero}. "
        )

        run.bold = True
        run.font.color.rgb = RGBColor(
            31,
            78,
            120
        )

        p.add_run(
            conclusion
        )

    # =====================================================
    # FIRMA
    # =====================================================

    doc.add_paragraph()
    doc.add_paragraph()

    firma = doc.add_paragraph(
        "____________________________________________"
    )

    firma.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()

    run = p.add_run(
        "Generado por: "
    )

    run.bold = True

    p.add_run(
        generado_por
    )

    doc.add_paragraph(
        "Responsable del Sistema de Gestión de Seguridad "
        "y Salud en el Trabajo"
    )

    # =====================================================
    # PIE DE PÁGINA
    # =====================================================

    pie = seccion.footer

    p_pie = pie.paragraphs[0]

    p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p_pie.add_run(
        "Informe de Gestión SG-SST"
    )

    run.font.size = Pt(7)
    run.font.color.rgb = GRIS_TEXTO

    # =====================================================
    # GUARDAR
    # =====================================================

    os.makedirs(
        Config.REPORTS_FOLDER,
        exist_ok=True
    )

    ruta = os.path.join(
        Config.REPORTS_FOLDER,
        nombre_archivo
    )

    doc.save(ruta)

    return ruta